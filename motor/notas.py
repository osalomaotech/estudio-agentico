"""Grava as notas do apresentador no .pptx a partir do campo "notas" de cada slide do slides.json.

Uso: python notas.py <projeto> <arquivo.pptx> [--curta]

Formato das notas (escrito pelo agente de narrativa):
- só a fala, um parágrafo por ideia, linha em branco entre parágrafos;
- **negrito** só no que deve ser enfatizado com a voz;
- parágrafo que sai na versão curta começa com "(corta na versão curta)";
- "(pausa)" sozinho numa linha marca pausa longa.
Na versão curta, os parágrafos marcados como cortáveis são removidos.
"""
import json
import pathlib
import re
import sys

from pptx import Presentation
from pptx.util import Pt

CORTA = "(corta na versão curta)"


def paragrafos(texto: str, curta: bool) -> list[str]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", texto.strip()) if p.strip()]
    if curta:
        paras = [p for p in paras if not p.startswith(CORTA)]
    return paras


def main() -> None:
    projeto = pathlib.Path(sys.argv[1])
    deck = pathlib.Path(sys.argv[2])
    curta = "--curta" in sys.argv
    spec = json.loads((projeto / "slides.json").read_text(encoding="utf-8"))
    por_id = {s.get("id", f"S{i + 1:02d}"): s for i, s in enumerate(spec["slides"])}
    ids = json.loads(deck.with_suffix(".ids.json").read_text())

    p = Presentation(str(deck))
    for slide, sid in zip(p.slides, ids):
        tf = slide.notes_slide.notes_text_frame
        tf.clear()
        linhas: list[str] = []
        for k, para in enumerate(paragrafos(por_id[sid].get("notas", ""), curta)):
            if k:
                linhas.append("")
            linhas += para.split("\n")
        for j, linha in enumerate(linhas or [""]):
            par = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            for k, pedaco in enumerate(linha.split("**")):
                if not pedaco:
                    continue
                r = par.add_run()
                r.text = pedaco
                r.font.size = Pt(16)
                r.font.bold = k % 2 == 1
                if pedaco.startswith("(corta") or pedaco.strip() == "(pausa)":
                    r.font.italic, r.font.size = True, Pt(12)
    p.save(str(deck))
    print(f"notas aplicadas: {len(ids)} slides")


if __name__ == "__main__":
    main()
